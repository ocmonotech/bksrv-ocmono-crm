"""
Parsers for common biometric / Tabulator-style Excel exports (samples: Attendance Record Report,
Schedule Information Report, Exception Statistic Report, Statistical Report of Attendance).
"""

from __future__ import annotations

import io
import json
import re
from calendar import monthrange
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

# --- helpers ---


def cell_as_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        return str(val).strip()
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    return str(val).strip()


def parse_day_header_number(val: Any) -> Optional[int]:
    s = cell_as_str(val)
    if not s or not s.isdigit():
        return None
    n = int(s)
    if 1 <= n <= 31:
        return n
    return None


def find_report_month(matrix: List[List[Any]]) -> Tuple[Optional[int], Optional[int]]:
    """Return (year, month) from first YYYY-MM-DD ~ YYYY-MM-DD in the sheet."""
    pat = re.compile(
        r"(\d{4})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{1,2})\s*[~\-–]\s*(\d{4})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{1,2})",
        re.I,
    )
    for row in matrix[:40]:
        for val in row:
            s = cell_as_str(val)
            m = pat.search(s)
            if m:
                return int(m.group(1)), int(m.group(2))
    # single date
    pat2 = re.compile(r"(\d{4})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{1,2})")
    for row in matrix[:40]:
        for val in row:
            s = cell_as_str(val)
            m = pat2.search(s)
            if m:
                return int(m.group(1)), int(m.group(2))
    return None, None


def find_day_column_map(matrix: List[List[Any]], max_scan: int = 120) -> Tuple[Optional[int], Dict[int, int]]:
    """
    Find a row where many cells are day numbers 1..31; return (row_index, col_index -> day).
    """
    best: Tuple[int, Dict[int, int]] = (-1, {})
    for i, row in enumerate(matrix[:max_scan]):
        mp: Dict[int, int] = {}
        for j, val in enumerate(row):
            d = parse_day_header_number(val)
            if d is not None:
                mp[j] = d
        if len(mp) >= 12 and len(mp) > len(best[1]):
            best = (i, mp)
    if best[0] < 0:
        return None, {}
    return best[0], best[1]


def parse_time_tokens(cell_val: Any) -> List[Tuple[int, int]]:
    """Extract all HH:MM-like tokens from one cell as (hour, minute)."""
    s = cell_as_str(cell_val)
    if not s:
        return []
    out: List[Tuple[int, int]] = []
    # Some machine exports concatenate times with no separator, e.g. "10:2919:12".
    # So we intentionally avoid word boundaries and just scan for HH:MM/HH:MM:SS substrings.
    for m in re.finditer(r"(\d{1,2}):(\d{2})(?::\d{2})?", s):
        hh = int(m.group(1))
        mm = int(m.group(2))
        if 0 <= hh <= 23 and 0 <= mm <= 59:
            out.append((hh, mm))
    return out


def row_joined_text(row: List[Any]) -> str:
    return " ".join(cell_as_str(c) for c in row if cell_as_str(c))


def parse_id_name_from_row(row: List[Any]) -> Tuple[Optional[str], Optional[str]]:
    """Extract machine user id and display name from 'ID: 2 Name: ... Dept:' style row."""
    line = row_joined_text(row)
    id_m = re.search(r"ID\s*[:\.]?\s*(\d+)", line, re.I)
    mid = id_m.group(1) if id_m else None

    # Fallback: some sheets split "ID:" and number into adjacent cells.
    if not mid:
        for i, c in enumerate(row):
            s = cell_as_str(c)
            if re.search(r"^\s*ID\s*[:\.]?\s*$", s, re.I):
                for j in range(i + 1, min(i + 4, len(row))):
                    cand = cell_as_str(row[j])
                    if re.fullmatch(r"\d+", cand):
                        mid = cand
                        break
                if mid:
                    break
    if not mid:
        return None, None

    name: Optional[str] = None
    name_m = re.search(r"Name\s*[:\.]?\s*(.+?)(?:\s+Dept|$)", line, re.I | re.DOTALL)
    if name_m:
        name = name_m.group(1).strip()

    # Fallback: extract from cells around "Name"
    if not name:
        for i, c in enumerate(row):
            s = cell_as_str(c)
            m = re.search(r"Name\s*[:\.]?\s*(.*)$", s, re.I)
            if not m:
                continue
            tail = (m.group(1) or "").strip()
            if tail and not re.search(r"\bDept\b", tail, re.I):
                name = tail
                break
            chunks: List[str] = []
            for j in range(i + 1, min(i + 6, len(row))):
                nxt = cell_as_str(row[j]).strip()
                if not nxt:
                    continue
                if re.search(r"\bDept\b", nxt, re.I):
                    break
                chunks.append(nxt)
                if len(chunks) >= 3:
                    break
            if chunks:
                name = " ".join(chunks).strip()
                break

    if name:
        name = re.sub(r"\s+", " ", name).strip()
    return mid, name or None


def row_has_times_in_day_cols(row: List[Any], col_to_day: Dict[int, int]) -> bool:
    for j in col_to_day:
        if j < len(row) and parse_time_tokens(row[j]):
            return True
    return False


def row_mostly_empty(row: List[Any]) -> bool:
    n = sum(1 for c in row if cell_as_str(c))
    return n <= 1


@dataclass
class ParsedMachineUpload:
    format_name: str
    report_year: Optional[int] = None
    report_month: Optional[int] = None
    punches: List[Tuple[str, datetime, str]] = field(default_factory=list)
    schedule_days: List[Dict[str, Any]] = field(default_factory=list)
    exception_days: List[Dict[str, Any]] = field(default_factory=list)
    statistics_rows: List[Dict[str, Any]] = field(default_factory=list)


def _matrix_text_blob(matrix: List[List[Any]], rows: int = 60, cols: int = 35) -> str:
    parts = []
    for i in range(min(rows, len(matrix))):
        for j in range(min(cols, len(matrix[i]) if matrix[i] else 0)):
            parts.append(cell_as_str(matrix[i][j]).lower())
    return " ".join(parts)


def detect_format(matrix: List[List[Any]]) -> Optional[str]:
    blob = _matrix_text_blob(matrix)
    if "attendance record report" in blob:
        return "attendance_record_grid"
    if "schedule information report" in blob:
        return "schedule_information"
    if "exception statistic" in blob:
        return "exception_statistic"
    if "statistical report of attendance" in blob:
        return "statistical_summary"
    return None


def parse_attendance_record_grid(matrix: List[List[Any]]) -> ParsedMachineUpload:
    out = ParsedMachineUpload("attendance_record_grid")
    y, m = find_report_month(matrix)
    if not y or not m:
        return out
    out.report_year, out.report_month = y, m
    last_day = monthrange(y, m)[1]
    hdr_row, col_to_day = find_day_column_map(matrix)
    if hdr_row is None or not col_to_day:
        return out

    r = hdr_row + 1
    n = len(matrix)
    while r < n:
        mid, nm = parse_id_name_from_row(matrix[r])
        if not mid:
            r += 1
            continue

        # Capture all timing rows for this employee block until next ID row.
        next_id_row = n
        for k in range(r + 1, n):
            next_mid, _ = parse_id_name_from_row(matrix[k])
            if next_mid:
                next_id_row = k
                break

        found_any = False
        for rr in range(r + 1, next_id_row):
            row_data = matrix[rr]
            if not row_has_times_in_day_cols(row_data, col_to_day):
                continue
            found_any = True
            for col_j, day in col_to_day.items():
                if day > last_day or col_j >= len(row_data):
                    continue
                for hh, mm in parse_time_tokens(row_data[col_j]):
                    try:
                        dt = datetime(y, m, day, hh, mm, 0)
                    except ValueError:
                        continue
                    note = json.dumps({"machine_id": mid, "name": nm, "day": day}, default=str)
                    out.punches.append((mid, dt, note))
                    # Also emit with parsed display name so name-first matching can resolve
                    # even when machine ID doesn't equal CRM users.id.
                    if nm:
                        out.punches.append((nm, dt, note))

        # Prevent infinite loops if malformed blocks are encountered.
        r = next_id_row if next_id_row > r else r + 1

    return out


def parse_schedule_information(matrix: List[List[Any]]) -> ParsedMachineUpload:
    out = ParsedMachineUpload("schedule_information")
    y, m = find_report_month(matrix)
    if not y or not m:
        return out
    out.report_year, out.report_month = y, m
    hdr_row, col_to_day = find_day_column_map(matrix)
    if hdr_row is None or not col_to_day:
        return out
    last_day = monthrange(y, m)[1]
    # data typically starts 2 rows below day numbers (weekday row in between)
    data_start = hdr_row + 2
    # infer fixed columns: smallest day column index -> id column is j-2 or 0
    first_day_col = min(col_to_day.keys())
    id_col = max(0, first_day_col - 2)
    name_col = max(0, first_day_col - 1)

    for r in range(data_start, len(matrix)):
        row = matrix[r]
        if id_col >= len(row):
            continue
        mid = cell_as_str(row[id_col])
        if not mid or not mid.isdigit():
            continue
        nm = cell_as_str(row[name_col]) if name_col < len(row) else ""
        for col_j, day in col_to_day.items():
            if day > last_day or col_j >= len(row):
                continue
            code = cell_as_str(row[col_j])
            if code.lower() in ("null", "none"):
                code = ""
            try:
                wd = date(y, m, day)
            except ValueError:
                continue
            out.schedule_days.append(
                {
                    "machine_user_key": mid,
                    "display_name": nm,
                    "work_date": wd.isoformat(),
                    "code": code or None,
                }
            )

    return out


def _find_header_row_with_tokens(matrix: List[List[Any]], tokens: List[str]) -> Optional[int]:
    tl = [t.lower() for t in tokens]
    for i, row in enumerate(matrix[:80]):
        joined = " ".join(cell_as_str(c).lower() for c in row)
        if all(t in joined for t in tl):
            return i
    return None


def _find_col(row: List[Any], *labels: str) -> Optional[int]:
    for j, val in enumerate(row):
        s = cell_as_str(val).lower()
        for lb in labels:
            if lb.lower() == s or lb.lower() in s:
                return j
    return None


def _parse_int_cell(val: Any) -> Optional[int]:
    s = cell_as_str(val)
    if not s:
        return None
    m = re.match(r"^-?\d+", s.replace(",", ""))
    if not m:
        return None
    try:
        return int(m.group(0))
    except ValueError:
        return None


def _parse_time_cell(val: Any) -> Optional[time]:
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, time):
        return val
    s = cell_as_str(val)
    if not s:
        return None
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$", s)
    if m:
        return time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))
    return None


def parse_exception_statistic(matrix: List[List[Any]]) -> ParsedMachineUpload:
    out = ParsedMachineUpload("exception_statistic")
    y, m = find_report_month(matrix)
    if not y or not m:
        return out
    out.report_year, out.report_month = y, m

    hdr = _find_header_row_with_tokens(matrix, ["date", "id"])
    if hdr is None:
        hdr = _find_header_row_with_tokens(matrix, ["date", "name"])
    if hdr is None:
        return out

    header = matrix[hdr]
    # scan next row for subheaders (On-duty / Off-duty repeated)
    header2 = matrix[hdr + 1] if hdr + 1 < len(matrix) else []
    col_date = _find_col(header, "date")
    col_id = _find_col(header, "id")
    col_name = _find_col(header, "name")
    if col_date is None:
        return out

    def find_grouped_duty_cols() -> List[Tuple[int, int]]:
        """Return list of (on_col, off_col) pairs left-to-right."""
        pairs: List[Tuple[int, int]] = []
        on_pat = re.compile(r"on[-\s]?duty", re.I)
        off_pat = re.compile(r"off[-\s]?duty", re.I)
        combined = [header, header2]
        for pass_i in range(2):
            row = combined[pass_i] if pass_i < len(combined) else []
            for j, val in enumerate(row):
                s = cell_as_str(val)
                if on_pat.search(s):
                    # find nearest off-duty to the right
                    for k in range(j + 1, min(j + 6, len(row))):
                        if off_pat.search(cell_as_str(row[k])):
                            pairs.append((j, k))
                            break
        # dedupe by first on col
        seen = set()
        uniq = []
        for p in pairs:
            if p[0] not in seen:
                seen.add(p[0])
                uniq.append(p)
        return uniq

    duty_pairs = find_grouped_duty_cols()
    # late / early / absence columns: search in header rows
    late_col = early_col = abs_col = total_col = None
    for scan in (header, header2):
        for j, val in enumerate(scan):
            s = cell_as_str(val).lower()
            if "late" in s and "min" in s:
                late_col = j
            if "leave" in s and "early" in s and "min" in s:
                early_col = j
            if "absence" in s and "min" in s:
                abs_col = j
            if s.strip() == "total(min)" or (s.startswith("total") and "min" in s):
                total_col = j

    data_start = hdr + 2
    if hdr + 1 < len(matrix) and duty_pairs:
        # if second row is subheader, data starts after it
        if any(cell_as_str(c) for c in header2):
            data_start = hdr + 2
        else:
            data_start = hdr + 1

    for r in range(data_start, len(matrix)):
        row = matrix[r]
        if col_date >= len(row):
            continue
        d_raw = row[col_date]
        if isinstance(d_raw, datetime):
            wd = d_raw.date()
        elif isinstance(d_raw, date):
            wd = d_raw
        else:
            ds = cell_as_str(d_raw)
            mdt = re.search(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", ds)
            if not mdt:
                continue
            wd = date(int(mdt.group(1)), int(mdt.group(2)), int(mdt.group(3)))

        mid = (
            cell_as_str(row[col_id])
            if col_id is not None and col_id < len(row)
            else ""
        )
        nm = (
            cell_as_str(row[col_name])
            if col_name is not None and col_name < len(row)
            else ""
        )
        if not mid and not nm:
            continue

        late_min = _parse_int_cell(row[late_col]) if late_col is not None and late_col < len(row) else None
        early_min = _parse_int_cell(row[early_col]) if early_col is not None and early_col < len(row) else None
        absence_min = _parse_int_cell(row[abs_col]) if abs_col is not None and abs_col < len(row) else None
        total_min = _parse_int_cell(row[total_col]) if total_col is not None and total_col < len(row) else None

        ex = {
            "machine_user_key": mid or nm,
            "display_name": nm,
            "work_date": wd.isoformat(),
            "late_min": late_min,
            "early_min": early_min,
            "absence_min": absence_min,
            "total_min": total_min,
        }
        out.exception_days.append(ex)

        key = mid or nm
        for on_c, off_c in duty_pairs:
            if on_c < len(row):
                t_on = _parse_time_cell(row[on_c])
                if t_on:
                    out.punches.append(
                        (
                            key,
                            datetime.combine(wd, t_on),
                            json.dumps({"source": "exception_on_duty"}, default=str),
                        )
                    )
            if off_c < len(row):
                t_off = _parse_time_cell(row[off_c])
                if t_off:
                    out.punches.append(
                        (
                            key,
                            datetime.combine(wd, t_off),
                            json.dumps({"source": "exception_off_duty"}, default=str),
                        )
                    )

    return out


def _merged_column_labels(matrix: List[List[Any]], id_row: int) -> List[str]:
    """Combine two header rows into one label per column (handles merged parent headers)."""
    r1 = matrix[id_row]
    r2 = matrix[id_row + 1] if id_row + 1 < len(matrix) else []
    n = max(len(r1), len(r2))
    tops: List[str] = []
    last_top = ""
    for j in range(n):
        top = cell_as_str(r1[j]).lower() if j < len(r1) else ""
        if top:
            last_top = top
        tops.append(last_top)
    labels: List[str] = []
    for j in range(n):
        bot = cell_as_str(r2[j]).lower() if j < len(r2) else ""
        lab = (tops[j] + " " + bot).strip()
        labels.append(lab)
    return labels


def parse_statistical_summary(matrix: List[List[Any]]) -> ParsedMachineUpload:
    out = ParsedMachineUpload("statistical_summary")
    y, m = find_report_month(matrix)
    if not y or not m:
        return out
    out.report_year, out.report_month = y, m

    id_row = None
    for i, row in enumerate(matrix[:25]):
        for j, val in enumerate(row):
            if cell_as_str(val).lower() == "id":
                id_row = i
                break
        if id_row is not None:
            break
    if id_row is None:
        return out

    header = matrix[id_row]
    labels = _merged_column_labels(matrix, id_row)

    def col_by_label(pred) -> Optional[int]:
        for j, lab in enumerate(labels):
            if pred(lab):
                return j
        return None

    col_id = col_by_label(lambda s: s == "id")
    col_name = col_by_label(lambda s: "name" in s and "department" not in s)
    late_times_j = col_by_label(lambda s: "late" in s and "times" in s)
    late_min_j = col_by_label(lambda s: "late" in s and "min" in s and "times" not in s)
    early_times_j = col_by_label(lambda s: ("leave" in s or "earl" in s) and "times" in s)
    early_min_j = col_by_label(
        lambda s: ("leave" in s or "earl" in s) and "min" in s and "times" not in s
    )
    absent_j = col_by_label(lambda s: "absent" in s)
    att_ratio_j = col_by_label(lambda s: "att." in s or "att " in s)

    if absent_j is None:
        for j, val in enumerate(header):
            if "absent" in cell_as_str(val).lower():
                absent_j = j
                break

    data_start = id_row + 2
    for r in range(data_start, len(matrix)):
        row = matrix[r]
        if col_id is None or col_id >= len(row):
            continue
        mid = cell_as_str(row[col_id])
        if not mid or not re.match(r"^\d+$", mid):
            continue
        name = cell_as_str(row[col_name]) if col_name is not None and col_name < len(row) else ""
        rec: Dict[str, Any] = {
            "machine_id": mid,
            "name": name,
            "late_times": _parse_int_cell(row[late_times_j]) if late_times_j is not None and late_times_j < len(row) else None,
            "late_min": _parse_int_cell(row[late_min_j]) if late_min_j is not None and late_min_j < len(row) else None,
            "early_times": _parse_int_cell(row[early_times_j]) if early_times_j is not None and early_times_j < len(row) else None,
            "early_min": _parse_int_cell(row[early_min_j]) if early_min_j is not None and early_min_j < len(row) else None,
            "absent_days": _parse_int_cell(row[absent_j]) if absent_j is not None and absent_j < len(row) else None,
            "att_days_ratio": cell_as_str(row[att_ratio_j]) if att_ratio_j is not None and att_ratio_j < len(row) else None,
        }
        out.statistics_rows.append(rec)

    return out


def parse_known_format(matrix: List[List[Any]], fmt: str) -> ParsedMachineUpload:
    if fmt == "attendance_record_grid":
        return parse_attendance_record_grid(matrix)
    if fmt == "schedule_information":
        return parse_schedule_information(matrix)
    if fmt == "exception_statistic":
        return parse_exception_statistic(matrix)
    if fmt == "statistical_summary":
        return parse_statistical_summary(matrix)
    return ParsedMachineUpload("unknown")


def load_matrix_xlsx(content: bytes) -> List[List[Any]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows: List[List[Any]] = []
    for row in ws.iter_rows(max_row=800, max_col=60, values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


def load_matrix_xls(content: bytes) -> List[List[Any]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sh = book.sheet_by_index(0)
    out: List[List[Any]] = []
    for ri in range(sh.nrows):
        r: List[Any] = []
        for ci in range(sh.ncols):
            c = sh.cell(ri, ci)
            if c.ctype == xlrd.XL_CELL_DATE:
                try:
                    from xlrd.xldate import xldate_as_datetime

                    r.append(xldate_as_datetime(c.value, book.datemode))
                except Exception:
                    r.append(c.value)
            else:
                r.append(c.value)
        out.append(r)
    return out


def try_parse_tabulator_workbook(content: bytes, filename: str) -> Optional[ParsedMachineUpload]:
    """Return parsed Tabulator export, or None if the workbook is not a recognized report layout."""
    lower = (filename or "").lower()
    if lower.endswith(".xlsx"):
        matrix = load_matrix_xlsx(content)
    elif lower.endswith(".xls"):
        matrix = load_matrix_xls(content)
    else:
        return None

    fmt = detect_format(matrix)
    if not fmt:
        return None
    return parse_known_format(matrix, fmt)
